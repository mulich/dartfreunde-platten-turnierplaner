"""Extract the original generator's workbooks without executing Excel formulas.

Requires openpyxl for extraction only; the web app imports the resulting JSON
using the Python standard library.
"""
import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl


def extract(path):
    match = re.fullmatch(r'Dartturnier (\d{4} \d{2} \d{2} \d{2}-\d{2})\.xlsx', path.name)
    if not match:
        raise ValueError(f'No source date in filename: {path.name}')
    original = datetime.strptime(match[1], '%Y %m %d %H-%M').replace(tzinfo=ZoneInfo('Europe/Berlin'))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = wb['Spielplan']
        if tuple(c.value for c in next(sheet.rows))[:7] != ('Spiel', 'Runde', 'Durchgang', 'Scheibe', 'Spieler 1', 'Spieler 2', 'Ergebnis'):
            raise ValueError(f'Unknown workbook layout: {path.name}')
        matches = []
        for row in sheet.iter_rows(min_row=2):
            values = [c.value for c in row]
            if all(v is None for v in values[:8]):
                continue
            if type(values[0]) is not int:
                raise ValueError(f'Invalid match row: {path.name}, {row[0].row}')
            matches.append(dict(zip(('number', 'round', 'wave', 'board', 'player1', 'player2', 'score1', 'score2'), values[:8]), source_row=row[0].row))
        # Hidden column L preserves original player order and tie ordering.
        players = [r[0] for r in wb['Tabelle'].iter_rows(min_row=2, min_col=12, max_col=12, values_only=True) if r[0] is not None]
        if not players or any(not isinstance(p, str) or p.startswith('=') for p in players):
            raise ValueError(f'Missing original player list: {path.name}')
        return {'name': f'Dartturnier · {original:%d.%m.%Y} · {original:%H:%M}',
                'created_at': original.astimezone(timezone.utc).isoformat(),
                'source_file': path.name, 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'source_notes': [], 'boards': 3 if any(m['board']=='Schwarz' for m in matches) else 2,
                'players': players, 'matches': matches}
    finally:
        wb.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('files', type=Path, nargs='+')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    payload = [extract(path) for path in args.files]
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
